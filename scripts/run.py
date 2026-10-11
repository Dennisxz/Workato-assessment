import csv
from openpyxl import load_workbook

input_file = "data/W_Recipe_Data.xlsx"
output_file = "data/W_Recipe_Data.csv"
sheet_name = "D_Data"

workbook = load_workbook(input_file, data_only=True)
sheet = workbook[sheet_name]

with open(output_file, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)

    for row in sheet.iter_rows(values_only=True):
        writer.writerow(["" if value is None else value for value in row])

print(f"Created {output_file}")
